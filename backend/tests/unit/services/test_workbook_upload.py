# Copyright (c) 2024-2026 Henry Wiechert (datafeta.io). SPDX-License-Identifier: AGPL-3.0-only
"""Tests for workbook uploads: staging, the sheet picker contract and file cleanup."""

import asyncio
import io
import json
import os
import zipfile
from types import SimpleNamespace

import openpyxl
import pytest
from fastapi import UploadFile

from backend.exceptions import InvalidInputError, ResourceNotFoundError
from backend.services.connection_service import ConnectionService
from backend.session_state import ConnectionStateManager

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_DETAILS = json.dumps({"type": "csv"})


def run_async(coro):
    return asyncio.run(coro)


def _make_service(tmp_path):
    request = SimpleNamespace(
        app=SimpleNamespace(state=SimpleNamespace(upload_root_dir=str(tmp_path / "uploads")))
    )
    return ConnectionService(state_manager=ConnectionStateManager(), request=request)


def _xlsx_bytes(sheets=None, hidden=()):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in (sheets or {
        "Orders": [["id", "qty"], [1, 3], [2, 5]],
        "Customers": [["id", "name"], [1, "ada"]],
        "Secret": [["x"], [1]],
    }).items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    for name in hidden or ("Secret",):
        if name in wb.sheetnames:
            wb[name].sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(filename, payload, content_type=XLSX_MIME):
    return UploadFile(filename=filename, file=io.BytesIO(payload), headers={"content-type": content_type})


def _session_files(tmp_path):
    session_dir = tmp_path / "uploads" / "s1"
    return sorted(p.name for p in session_dir.iterdir()) if session_dir.exists() else []


def _staged(*refs):
    return json.dumps([{"upload_id": upload_id, "parts": parts} for upload_id, parts in refs])


class TestStageFiles:
    def test_lists_workbook_sheets(self, tmp_path):
        service = _make_service(tmp_path)

        result = run_async(service.stage_files([_upload("shop.xlsx", _xlsx_bytes())], "s1"))

        [entry] = result["uploads"]
        assert entry["filename"] == "shop.xlsx"
        assert (entry["format"], entry["part_label"]) == ("workbook", "sheet")
        assert entry["parts"] == [
            {"name": "Orders", "selectable": True, "reason": None},
            {"name": "Customers", "selectable": True, "reason": None},
            {"name": "Secret", "selectable": False, "reason": "hidden"},
        ]
        assert entry["upload_id"] in service.state_manager.staged_uploads

    def test_non_workbook_files_are_staged_without_sheets(self, tmp_path):
        service = _make_service(tmp_path)

        result = run_async(service.stage_files([_upload("a.csv", b"id\n1\n", "text/csv")], "s1"))

        assert result["uploads"][0]["format"] == "csv"
        assert result["uploads"][0]["parts"] is None

    def test_zip_members_staged_individually(self, tmp_path):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as archive:
            archive.writestr("shop.xlsx", _xlsx_bytes())
            archive.writestr("extra.csv", "id\n1\n")
        service = _make_service(tmp_path)

        result = run_async(service.stage_files([_upload("bundle.zip", buf.getvalue(), "application/zip")], "s1"))

        formats = {e["filename"]: e["format"] for e in result["uploads"]}
        assert formats == {"shop.xlsx": "workbook", "extra.csv": "csv"}

    def test_invalid_workbook_rejected_and_removed(self, tmp_path):
        service = _make_service(tmp_path)

        with pytest.raises(InvalidInputError):
            run_async(service.stage_files([_upload("bad.xlsx", b"id,name\n1,ada\n")], "s1"))

        assert _session_files(tmp_path) == []
        assert service.state_manager.staged_uploads == {}

    def test_staging_again_discards_previous_uploads(self, tmp_path):
        service = _make_service(tmp_path)
        first = run_async(service.stage_files([_upload("a.xlsx", _xlsx_bytes())], "s1"))
        run_async(service.stage_files([_upload("b.xlsx", _xlsx_bytes())], "s1"))

        assert first["uploads"][0]["upload_id"] not in service.state_manager.staged_uploads
        assert len(_session_files(tmp_path)) == 1

    def test_discard_staged_deletes_files(self, tmp_path):
        service = _make_service(tmp_path)
        run_async(service.stage_files([_upload("a.xlsx", _xlsx_bytes())], "s1"))

        run_async(service.discard_staged())

        assert service.state_manager.staged_uploads == {}
        assert _session_files(tmp_path) == []


class TestConnectWithStagedUploads:
    def _stage(self, service):
        result = run_async(service.stage_files([_upload("shop.xlsx", _xlsx_bytes())], "s1"))
        return result["uploads"][0]["upload_id"]

    def test_connect_loads_selected_sheets(self, tmp_path):
        service = _make_service(tmp_path)
        upload_id = self._stage(service)

        run_async(service.connect_multipart(CSV_DETAILS, [], "s1", _staged((upload_id, ["Orders"]))))

        connector = service.state_manager.current_connector
        assert [t.name for t in connector.list_tables()] == ["shop_orders"]
        assert service.state_manager.staged_uploads == {}
        # The session tracks the upload; the connector owns the derived sheet Parquet.
        assert len(service.state_manager.current_temp_paths) == 1

    def test_omitted_sheets_load_all_visible(self, tmp_path):
        service = _make_service(tmp_path)
        upload_id = self._stage(service)

        run_async(service.connect_multipart(CSV_DETAILS, [], "s1", json.dumps([{"upload_id": upload_id}])))

        tables = [t.name for t in service.state_manager.current_connector.list_tables()]
        assert tables == ["shop_orders", "shop_customers"]

    def test_staged_upload_consumed_once(self, tmp_path):
        service = _make_service(tmp_path)
        upload_id = self._stage(service)
        run_async(service.connect_multipart(CSV_DETAILS, [], "s1", _staged((upload_id, ["Orders"]))))

        with pytest.raises(ResourceNotFoundError):
            run_async(service.connect_multipart(CSV_DETAILS, [], "s1", _staged((upload_id, ["Orders"]))))

    def test_invalid_sheet_keeps_staged_upload_for_retry(self, tmp_path):
        service = _make_service(tmp_path)
        upload_id = self._stage(service)

        with pytest.raises(InvalidInputError, match="Secret"):
            run_async(service.connect_multipart(CSV_DETAILS, [], "s1", _staged((upload_id, ["Secret"]))))

        assert upload_id in service.state_manager.staged_uploads
        run_async(service.connect_multipart(CSV_DETAILS, [], "s1", _staged((upload_id, ["Customers"]))))

    def test_malformed_staged_json_rejected(self, tmp_path):
        service = _make_service(tmp_path)

        with pytest.raises(InvalidInputError) as excinfo:
            run_async(service.connect_multipart(CSV_DETAILS, [], "s1", '{"upload_id": "x"}'))
        assert excinfo.value.status_code == 422

    def test_direct_upload_applies_saved_selection(self, tmp_path):
        service = _make_service(tmp_path)
        details = json.dumps({"type": "csv", "file_parts": {"shop.xlsx": ["Customers"]}})

        run_async(service.connect_multipart(details, [_upload("shop.xlsx", _xlsx_bytes())], "s1"))

        tables = [t.name for t in service.state_manager.current_connector.list_tables()]
        assert tables == ["shop_customers"]

    def test_direct_upload_with_missing_saved_sheet_fails_cleanly(self, tmp_path):
        service = _make_service(tmp_path)
        details = json.dumps({"type": "csv", "file_parts": {"shop.xlsx": ["Archive"]}})

        with pytest.raises(InvalidInputError, match="'Archive'"):
            run_async(service.connect_multipart(details, [_upload("shop.xlsx", _xlsx_bytes())], "s1"))

        assert _session_files(tmp_path) == []
        assert service.state_manager.current_connector is None

    def test_sqlite_rejects_staged_uploads(self, tmp_path):
        service = _make_service(tmp_path)
        upload_id = self._stage(service)

        with pytest.raises(InvalidInputError, match="Staged uploads"):
            run_async(service.connect_multipart(
                json.dumps({"type": "sqlite"}), [], "s1", _staged((upload_id, None))
            ))


class TestAddFilesWithStagedUploads:
    def test_add_workbook_sheets(self, tmp_path):
        service = _make_service(tmp_path)
        run_async(service.connect_multipart(CSV_DETAILS, [_upload("base.csv", b"id\n1\n", "text/csv")], "s1"))
        staged = run_async(service.stage_files([_upload("shop.xlsx", _xlsx_bytes())], "s1"))
        upload_id = staged["uploads"][0]["upload_id"]

        result = run_async(service.add_files([], "s1", _staged((upload_id, ["Orders", "Customers"]))))

        assert result["added_tables"] == ["shop_orders", "shop_customers"]
        assert result["skipped_parts"] == []
        # The uploads base.csv and shop.xlsx; the sheet Parquet belongs to the connector.
        assert len(service.state_manager.current_temp_paths) == 2

    def test_add_reports_skipped_empty_sheet(self, tmp_path):
        service = _make_service(tmp_path)
        run_async(service.connect_multipart(CSV_DETAILS, [_upload("base.csv", b"id\n1\n", "text/csv")], "s1"))
        payload = _xlsx_bytes({"Data": [["a"], [1]], "Blank": []}, hidden=())
        staged = run_async(service.stage_files([_upload("more.xlsx", payload)], "s1"))

        result = run_async(service.add_files([], "s1", _staged((staged["uploads"][0]["upload_id"], None))))

        assert result["added_tables"] == ["more_data"]
        assert result["skipped_parts"] == ["more.xlsx: Blank"]


class TestDisconnectCleanup:
    def test_disconnect_removes_workbook_sheet_parquet_and_staged_files(self, tmp_path):
        service = _make_service(tmp_path)
        run_async(service.connect_multipart(CSV_DETAILS, [_upload("shop.xlsx", _xlsx_bytes())], "s1"))
        run_async(service.stage_files([_upload("pending.xlsx", _xlsx_bytes())], "s1"))
        assert len(_session_files(tmp_path)) == 4  # workbook, 2 sheet Parquet, staged workbook

        run_async(service.disconnect("s1"))

        assert _session_files(tmp_path) == []

    def test_disconnect_removes_flattened_json_parquet(self, tmp_path):
        """Regression: JSON's derived __flat.parquet used to survive until shutdown."""
        service = _make_service(tmp_path)
        payload = b'[{"id": 1}, {"id": 2}]'
        run_async(service.connect_multipart(CSV_DETAILS, [_upload("e.json", payload, "application/json")], "s1"))
        assert any(name.endswith("__flat.parquet") for name in _session_files(tmp_path))

        run_async(service.disconnect("s1"))

        assert _session_files(tmp_path) == []

    def test_reconnect_removes_previous_derived_files(self, tmp_path):
        service = _make_service(tmp_path)
        run_async(service.connect_multipart(CSV_DETAILS, [_upload("shop.xlsx", _xlsx_bytes())], "s1"))

        run_async(service.connect_multipart(CSV_DETAILS, [_upload("a.csv", b"id\n1\n", "text/csv")], "s1"))

        assert len(_session_files(tmp_path)) == 1
        assert not any(name.endswith(".parquet") for name in _session_files(tmp_path))
