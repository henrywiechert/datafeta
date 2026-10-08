# Copyright (c) 2024-2026 Henry Wiechert (datafeta.io). SPDX-License-Identifier: AGPL-3.0-only
"""Contract every registered file format handler must meet.

FileConnector, ConnectionService and the frontend rely only on this
contract, so a new format is complete once it is registered and passes here.
Registering a format without adding a sample below fails the suite.
"""

import os

import duckdb
import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
import pytest

from backend.connectors.file_handlers import (
    CSV_OPTIONS,
    DATE_FORMAT_OPTIONS,
    FILE_EXTENSIONS,
    FILE_HANDLERS,
    OpenedFile,
    file_format_catalog,
    handler_for,
)
from backend.exceptions import InvalidInputError
from backend.routers.app_config import get_app_config

KNOWN_OPTION_GROUPS = {CSV_OPTIONS, DATE_FORMAT_OPTIONS}


def _csv(path):
    path.write_text("id,name\n1,ada\n2,bob\n")


def _parquet(path):
    pq.write_table(pa.table({"id": [1, 2], "name": ["ada", "bob"]}), path)


def _json(path):
    path.write_text('[{"id": 1, "name": "ada"}, {"id": 2, "name": "bob"}]')


def _workbook(path):
    wb = openpyxl.Workbook()
    wb.active.title = "People"
    for row in (["id", "name"], [1, "ada"], [2, "bob"]):
        wb.active.append(row)
    wb.save(path)


# One minimal two-row file per registered format key.
SAMPLES = {
    "csv": (".csv", _csv),
    "parquet": (".parquet", _parquet),
    "json": (".json", _json),
    "workbook": (".xlsx", _workbook),
}


def _sample(tmp_path, handler_cls):
    extension, write = SAMPLES[handler_cls.FORMAT.key]
    path = tmp_path / f"sample{extension}"
    write(path)
    return str(path)


def test_every_format_has_a_contract_sample():
    assert {h.FORMAT.key for h in FILE_HANDLERS} == set(SAMPLES)


def test_format_declarations_are_unique_and_well_formed():
    keys = [h.FORMAT.key for h in FILE_HANDLERS]
    extensions = [ext for h in FILE_HANDLERS for ext in h.FORMAT.extensions]
    assert len(keys) == len(set(keys))
    assert len(extensions) == len(set(extensions)), "an extension maps to one handler"
    for ext in extensions:
        assert ext.startswith(".") and ext == ext.lower()
    for handler in FILE_HANDLERS:
        assert handler.FORMAT.options <= KNOWN_OPTION_GROUPS
        assert handler.FORMAT.label


def test_catalog_is_published_to_the_frontend():
    assert get_app_config()["fileFormats"] == file_format_catalog()
    assert {ext for f in file_format_catalog() for ext in f["extensions"]} == FILE_EXTENSIONS


def test_unknown_extension_rejected():
    with pytest.raises(InvalidInputError, match="Unsupported file type: .txt"):
        handler_for("notes.txt")


@pytest.mark.parametrize("handler_cls", FILE_HANDLERS, ids=lambda h: h.FORMAT.key)
class TestHandlerContract:
    def test_handler_for_resolves_every_extension(self, handler_cls):
        for ext in handler_cls.FORMAT.extensions:
            assert type(handler_for(f"file{ext.upper()}")) is handler_cls

    def test_validates_sample_and_rejects_garbage(self, handler_cls, tmp_path):
        handler = handler_for(_sample(tmp_path, handler_cls))
        handler.validate(_sample(tmp_path, handler_cls))

        garbage = tmp_path / f"garbage{handler_cls.FORMAT.extensions[0]}"
        garbage.write_bytes(b"")
        with pytest.raises(InvalidInputError):
            handler.validate(str(garbage))

    def test_opens_into_readable_tables(self, handler_cls, tmp_path):
        path = _sample(tmp_path, handler_cls)
        handler = handler_for(path, {})

        opened = handler.open(path, None, os.path.basename(path))

        assert isinstance(opened, OpenedFile) and opened.tables
        con = duckdb.connect()
        for table in opened.tables:
            rows = con.execute(f"SELECT * FROM {table.reader.build_reader_sql(table.path)}").fetchall()
            assert len(rows) == 2
            # Derived files live next to the upload and are flagged for deletion.
            assert table.generated == (table.path != path)
            if table.generated:
                assert os.path.dirname(table.path) == os.path.dirname(path)

    def test_parts_match_part_label(self, handler_cls, tmp_path):
        parts = handler_for(_sample(tmp_path, handler_cls)).list_parts(_sample(tmp_path, handler_cls))
        if handler_cls.FORMAT.part_label is None:
            assert parts is None
        else:
            assert parts and all(p.name for p in parts)
